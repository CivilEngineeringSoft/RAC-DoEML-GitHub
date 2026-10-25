"""Variables, specimen coding, strength conversion factors and mix construction.

Conventions of the paper (Section 2.1.1):
* Specimen coding: 100x200 mm cylinder = 0.1, 150x300 mm cylinder = 0.2,
  100 mm cube = 0.3, 150 mm cube = 0.4, 75x150 mm cylinder = 0.5;
* Strength classification feature (SCF): 1 = high strength (expected fc >= 45 MPa),
  0 = low strength (expected fc < 45 MPa);
* Strength (specimen shape) coefficient: sigma_cyl / sigma_x from the conversion
  table (shape, size and strength grade), see :data:`CONVERSION`;
* an empty cell means that the constituent is absent and is replaced by 0
  (this reproduces Table 2 of the paper).
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np
import pandas as pd

APP_DIR = Path(__file__).resolve().parents[1]
DATA_FILE = APP_DIR / "data" / "RAC_28d_strength_dataset.xlsx"

TARGET = "CS (MPa)"
SCF = "Strength classification feature (Expected_strength)"
COEF = "Strength coefficient"

# 20 model inputs, in the order of the dataset
FEATURES: List[str] = [
    "Specimen",
    "Size RCA (mm)", "Absorption RCA (%)", "Quantity RCA (kg/m3)", "SSD RCA (kg/m3)",
    "Absorption RFA (%)", "RFA content (kg/m3)",
    "Size NCA (mm)", "Absorption NCA (%)", "Quantity NCA (kg/m3)", "SSD NCA (kg/m3)",
    "Quantity FNA (kg/m3)",
    "Coarse aggregate replacement ratio (%)", "Fine aggregate replacement ratio (%)",
    "Fly Ash Content (kg/m3)", "Water-to-binder ratio (W/B)", "Cement Content (kg/m3)",
    "Admixture Content (%)",
    SCF, COEF,
]

# ------------------------------------------------------------------ specimens
SPECIMENS: Dict[str, dict] = {
    # key: code used by the models, conversion-table row, label
    "100x200": {"code": 0.1, "row": ("Cylinder", "100x200"), "label": "Cylinder 100 x 200 mm"},
    "150x300": {"code": 0.2, "row": ("Cylinder", "150x300"), "label": "Cylinder 150 x 300 mm"},
    "100x100x100": {"code": 0.3, "row": ("Cube", "100"), "label": "Cube 100 mm"},
    "150x150x150": {"code": 0.4, "row": ("Cube", "150"), "label": "Cube 150 mm"},
    "75x150": {"code": 0.5, "row": ("Cylinder", "75x150"), "label": "Cylinder 75 x 150 mm"},
}

GRADES = ["C20-40", "C50", "C60", "C70", "C80"]

# Conversion factors sigma_cyl / sigma_x towards the 160 x 320 mm cylinder
# (Table "Conversion Factors for Compressive Strength Across Specimen Shapes and Sizes")
CONVERSION: Dict[tuple, List[float]] = {
    ("Cylinder", "160x320"): [1.00, 1.00, 1.00, 1.00, 1.00],
    ("Cylinder", "50x100"): [0.93, 0.93, 0.88, 0.88, 0.88],
    ("Cylinder", "75x150"): [0.94, 0.95, 0.91, 0.91, 0.92],
    ("Cylinder", "100x200"): [0.96, 0.97, 0.96, 0.96, 0.96],
    ("Cylinder", "120x240"): [0.98, 0.98, 0.97, 0.97, 0.97],
    ("Cylinder", "150x300"): [0.99, 0.99, 0.99, 0.99, 0.99],
    ("Prism", "100x300"): [0.98, 1.02, 1.01, 1.00, 1.00],
    ("Prism", "120x360"): [1.00, 1.04, 1.04, 1.03, 1.03],
    ("Prism", "150x300"): [1.05, 1.09, 1.09, 1.09, 1.08],
    ("Cube", "150"): [0.79, 0.83, 0.85, 0.87, 0.88],
    ("Cube", "100"): [0.75, 0.79, 0.81, 0.83, 0.84],
}


def conversion_table() -> pd.DataFrame:
    rows = [{"Shape": s, "Size (mm)": z, **dict(zip(GRADES, v))} for (s, z), v in CONVERSION.items()]
    return pd.DataFrame(rows)


def grade_from_cylinder_strength(fc_cyl: float) -> str:
    """Strength grade from the equivalent 160x320 cylinder strength (MPa).

    C20-40 below 45 MPa, then classes centred on 50, 60, 70 and 80 MPa
    (limits 45, 55, 65 and 75 MPa)."""
    for lim, g in ((45, "C20-40"), (55, "C50"), (65, "C60"), (75, "C70")):
        if fc_cyl < lim:
            return g
    return "C80"


def strength_coefficient(specimen: str, grade: str) -> float:
    """Conversion factor sigma_cyl / sigma_specimen for a specimen key and a grade."""
    return CONVERSION[SPECIMENS[specimen]["row"]][GRADES.index(grade)]


def initial_grade(high_strength: bool) -> str:
    return "C50" if high_strength else "C20-40"


def specimen_key(label) -> Optional[str]:
    s = str(label).lower().replace("×", "x").replace("*", "x")
    s = re.sub(r"mm|\s|cylinder|cyl|cube", "", s)
    if s in SPECIMENS:
        return s
    if s in ("100", "150"):
        return f"{s}x{s}x{s}"
    try:
        v = round(float(s), 2)
        return next((k for k, d in SPECIMENS.items() if d["code"] == v), None)
    except ValueError:
        return None


def specimen_code(label) -> float:
    k = specimen_key(label)
    if k is None:
        raise ValueError(f"Unknown specimen '{label}'. Allowed: "
                         + ", ".join(d["label"] for d in SPECIMENS.values()))
    return SPECIMENS[k]["code"]


# ------------------------------------------------------------------ dataset
def load_data() -> pd.DataFrame:
    df = pd.read_excel(DATA_FILE)
    df.columns = [re.sub(r"\s+", " ", str(c).strip()) for c in df.columns]
    return df


def prepare(df: pd.DataFrame) -> pd.DataFrame:
    """Numeric model matrix: specimen coding, empty cells -> 0, column order of FEATURES."""
    X = df.copy()
    X["Specimen"] = X["Specimen"].map(specimen_code)
    X = X.reindex(columns=FEATURES)
    return X.apply(pd.to_numeric, errors="coerce").fillna(0.0).astype(float)


# ------------------------------------------------------------------ mixes
def build_mix(*, specimen: str, high_strength: bool,
              cement: float, fly_ash: float, wb: float, admixture: float,
              fine_total: float, coarse_total: float,
              ca_ratio: float, fa_ratio: float,
              nca_size: float, nca_abs: float, nca_ssd: float,
              rca_size: float, rca_abs: float, rca_ssd: float,
              rfa_abs: float) -> Dict[str, float]:
    """Raw model inputs (without the strength coefficient) of one mix.

    Replacement ratios are by mass, as in the dataset: RCA = ratio x total coarse
    aggregate, RFA = ratio x total fine aggregate. Properties of an absent
    constituent are set to 0, as in the dataset."""
    rca = coarse_total * ca_ratio / 100.0
    nca = coarse_total - rca
    rfa = fine_total * fa_ratio / 100.0
    fna = fine_total - rfa
    return {
        "Specimen": SPECIMENS[specimen]["code"],
        "Size RCA (mm)": rca_size if rca > 0 else 0.0,
        "Absorption RCA (%)": rca_abs if rca > 0 else 0.0,
        "Quantity RCA (kg/m3)": rca,
        "SSD RCA (kg/m3)": rca_ssd if rca > 0 else 0.0,
        "Absorption RFA (%)": rfa_abs if rfa > 0 else 0.0,
        "RFA content (kg/m3)": rfa,
        "Size NCA (mm)": nca_size if nca > 0 else 0.0,
        "Absorption NCA (%)": nca_abs if nca > 0 else 0.0,
        "Quantity NCA (kg/m3)": nca,
        "SSD NCA (kg/m3)": nca_ssd if nca > 0 else 0.0,
        "Quantity FNA (kg/m3)": fna,
        "Coarse aggregate replacement ratio (%)": ca_ratio,
        "Fine aggregate replacement ratio (%)": fa_ratio,
        "Fly Ash Content (kg/m3)": fly_ash,
        "Water-to-binder ratio (W/B)": wb,
        "Cement Content (kg/m3)": cement,
        "Admixture Content (%)": admixture,
        SCF: 1.0 if high_strength else 0.0,
    }


def water_content(mix: Dict[str, float]) -> float:
    return mix["Water-to-binder ratio (W/B)"] * (mix["Cement Content (kg/m3)"] + mix["Fly Ash Content (kg/m3)"])


# ------------------------------------------------------------------ templates
TEMPLATE_COLS = ["Mix ID", "Specimen", "Expected strength (High/Low)", "Strength grade (optional)",
                 "Cement Content (kg/m3)", "Fly Ash Content (kg/m3)", "Water-to-binder ratio (W/B)",
                 "Admixture Content (%)",
                 "Quantity NCA (kg/m3)", "Size NCA (mm)", "Absorption NCA (%)", "SSD NCA (kg/m3)",
                 "Quantity RCA (kg/m3)", "Size RCA (mm)", "Absorption RCA (%)", "SSD RCA (kg/m3)",
                 "Quantity FNA (kg/m3)", "RFA content (kg/m3)", "Absorption RFA (%)"]


def mixes_template() -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    ex = pd.DataFrame([
        ["M1", "Cube 100 mm", "High", "", 400, 0, 0.50, 1.5, 1027.58, 10, 0.75, 2660, 0, "", "", "", 786.24, 0, ""],
        ["M2", "Cube 100 mm", "High", "", 400, 0, 0.50, 1.5, 822.06, 10, 0.75, 2660, 205.52, 10, 6.6, 2400,
         628.99, 157.25, 13.3],
    ], columns=TEMPLATE_COLS)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        ex.to_excel(xl, sheet_name="Mixes", index=False)
        ws = xl.sheets["Mixes"]
        for j in range(1, len(TEMPLATE_COLS) + 1):
            c = ws.cell(row=1, column=j)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            ws.column_dimensions[c.column_letter].width = 15
        ws.row_dimensions[1].height = 45
        for col, opts in (("B", ",".join(d["label"] for d in SPECIMENS.values())), ("C", "High,Low"),
                          ("D", ",".join(GRADES))):
            dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}2000")
        info = pd.DataFrame({"Notes": [
            "One row per mix. Quantities in kg/m3 of concrete; admixture in % of cement mass.",
            "Leave the properties of an absent constituent empty (e.g. no RCA -> RCA size, absorption and SSD empty).",
            "Replacement ratios are computed from the quantities: RCA/(RCA+NCA) and RFA/(RFA+FNA).",
            "Expected strength: High = expected fc >= 45 MPa, Low = expected fc < 45 MPa.",
            "Strength grade (optional): if empty, it is determined from the predicted strength.",
        ]})
        info.to_excel(xl, sheet_name="Instructions", index=False)
        xl.sheets["Instructions"].column_dimensions["A"].width = 110
    return buf.getvalue()


def read_mixes(uploaded) -> pd.DataFrame:
    name = getattr(uploaded, "name", str(uploaded)).lower()
    df = pd.read_csv(uploaded, sep=None, engine="python") if name.endswith(".csv") else pd.read_excel(uploaded)
    df.columns = [re.sub(r"\s+", " ", str(c).strip()) for c in df.columns]
    return df.dropna(how="all")


def mixes_from_table(df: pd.DataFrame) -> List[dict]:
    """Rows of the template -> list of dicts with raw inputs, specimen key, class and grade."""
    out = []
    num = lambda r, c: float(pd.to_numeric(r.get(c), errors="coerce")) if pd.notna(pd.to_numeric(r.get(c), errors="coerce")) else 0.0
    for i, r in df.iterrows():
        spec = specimen_key(r.get("Specimen"))
        if spec is None:
            raise ValueError(f"Row {i + 2}: unknown specimen '{r.get('Specimen')}'.")
        hs = str(r.get("Expected strength (High/Low)", "High")).strip().lower().startswith("h")
        rca, nca, rfa, fna = (num(r, "Quantity RCA (kg/m3)"), num(r, "Quantity NCA (kg/m3)"),
                              num(r, "RFA content (kg/m3)"), num(r, "Quantity FNA (kg/m3)"))
        mix = build_mix(
            specimen=spec, high_strength=hs, cement=num(r, "Cement Content (kg/m3)"),
            fly_ash=num(r, "Fly Ash Content (kg/m3)"), wb=num(r, "Water-to-binder ratio (W/B)"),
            admixture=num(r, "Admixture Content (%)"), fine_total=rfa + fna, coarse_total=rca + nca,
            ca_ratio=100 * rca / (rca + nca) if rca + nca > 0 else 0.0,
            fa_ratio=100 * rfa / (rfa + fna) if rfa + fna > 0 else 0.0,
            nca_size=num(r, "Size NCA (mm)"), nca_abs=num(r, "Absorption NCA (%)"), nca_ssd=num(r, "SSD NCA (kg/m3)"),
            rca_size=num(r, "Size RCA (mm)"), rca_abs=num(r, "Absorption RCA (%)"), rca_ssd=num(r, "SSD RCA (kg/m3)"),
            rfa_abs=num(r, "Absorption RFA (%)"))
        g = str(r.get("Strength grade (optional)", "") or "").strip().upper().replace(" ", "")
        grade = next((x for x in GRADES if x.upper() == g), None)
        out.append({"id": str(r.get("Mix ID", f"Mix {i + 1}")), "mix": mix, "specimen": spec,
                    "high": hs, "grade": grade})
    return out
